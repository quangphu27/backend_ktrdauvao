"""Xuất Excel kết quả quiz theo ID."""
from __future__ import annotations

import os
import sys
from io import BytesIO

from bson import ObjectId
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pymongo import MongoClient

load_dotenv()

QUIZ_ID = sys.argv[1] if len(sys.argv) > 1 else "6ab7e539c4110429822a0cb0"
OUT_DIR = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), "..", "frontend", "public")
FRONTEND = os.getenv("FRONTEND_URL") or "https://kiemtradauvao.vercel.app"


def oid_str(v):
    return str(v)


def main():
    client = MongoClient(os.getenv("MONGODB_URI"))
    db = client[os.getenv("MONGODB_DB_NAME")]

    try:
        quiz = db.quizzes.find_one({"_id": ObjectId(QUIZ_ID)})
    except Exception:
        quiz = None
    if not quiz:
        quiz = db.quizzes.find_one({"_id": QUIZ_ID})
    if not quiz:
        print("NOT_FOUND")
        sys.exit(1)

    quiz_oid = oid_str(quiz["_id"])
    attempts = list(
        db.quiz_attempts.find({"quiz_id": quiz_oid}).sort("submitted_at", -1)
    )
    frontend = FRONTEND.rstrip("/")

    def sheet_name(raw):
        name = (raw or "").strip() or "Khong_ro_lop"
        for ch in r"\/?*[]:":
            name = name.replace(ch, "_")
        return name[:31] or "Khong_ro_lop"

    def sort_key(grade):
        g = (grade or "").strip()
        order = {
            "Scratch1": 0,
            "Scratch2": 1,
            "Scratch3": 2,
            "Scratch4": 3,
        }
        if g in order:
            return (0, order[g], g.lower())
        if not g or g == "Không rõ lớp":
            return (2, 99, "")
        return (1, 0, g.lower())

    by_grade = {}
    for a in attempts:
        grade = (a.get("student_grade") or "").strip() or "Không rõ lớp"
        by_grade.setdefault(grade, []).append(a)
    grades_sorted = sorted(by_grade.keys(), key=sort_key)

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="4A90D9")
    header_font = Font(bold=True, color="FFFFFF")
    link_font = Font(color="0563C1", underline="single")
    headers = [
        "STT",
        "Ten hoc sinh",
        "Lop",
        "Diem %",
        "Diem (dat/tong)",
        "Link chi tiet bai lam",
        "Link admin cham bai",
        "Ngay nop",
    ]

    used = set()

    def unique_title(base):
        title = base
        n = 2
        while title in used:
            suffix = f"_{n}"
            title = base[: 31 - len(suffix)] + suffix
            n += 1
        used.add(title)
        return title

    def fill_sheet(ws, rows):
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
        for i, a in enumerate(rows, 1):
            attempt_id = oid_str(a["_id"])
            detail_url = f"{frontend}/quiz-ket-qua/{attempt_id}"
            admin_url = f"{frontend}/admin/quizzes/{quiz_oid}/results?attempt={attempt_id}"
            submitted = a.get("submitted_at")
            date_str = (
                submitted.strftime("%d/%m/%Y %H:%M")
                if hasattr(submitted, "strftime")
                else str(submitted or "")
            )
            ws.append(
                [
                    i,
                    a.get("student_name") or "",
                    a.get("student_grade") or "",
                    a.get("score"),
                    f"{a.get('earned')}/{a.get('max_score')}",
                    "Xem chi tiet",
                    "Cham bai",
                    date_str,
                ]
            )
            row = ws.max_row
            c_detail = ws.cell(row, 6)
            c_detail.hyperlink = detail_url
            c_detail.font = link_font
            c_admin = ws.cell(row, 7)
            c_admin.hyperlink = admin_url
            c_admin.font = link_font
        widths = {
            "A": 6,
            "B": 28,
            "C": 14,
            "D": 10,
            "E": 14,
            "F": 18,
            "G": 18,
            "H": 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

    fill_sheet(wb.create_sheet(unique_title("Tat_ca"), 0), attempts)
    for grade in grades_sorted:
        rows = sorted(by_grade[grade], key=lambda x: (x.get("student_name") or "").lower())
        fill_sheet(wb.create_sheet(unique_title(sheet_name(grade))), rows)

    raw_title = quiz.get("title") or "quiz"
    safe_title = (
        raw_title.encode("ascii", "ignore").decode("ascii")
        or "quiz"
    )
    safe_title = "".join(c if c.isalnum() or c in " _-" else "_" for c in safe_title)[:40].strip().replace(" ", "_") or "quiz"
    out_path = os.path.abspath(os.path.join(OUT_DIR, f"ket_qua_{safe_title}_{QUIZ_ID[:8]}.xlsx"))
    # luôn có bản tên ASCII rõ ràng
    out_path = os.path.abspath(os.path.join(OUT_DIR, f"ket_qua_quiz_{QUIZ_ID[:8]}.xlsx"))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print(f"OK title={quiz.get('title')}")
    print(f"attempts={len(attempts)}")
    print(f"grades={len(grades_sorted)}")
    print(f"file={out_path}")


if __name__ == "__main__":
    main()
