"""Cập nhật 4 file phiếu nhận xét Scratch:
- Cột Điểm = chỉ điểm trắc nghiệm (/100)
- Cột «Xem bài làm» (bên phải Nhận xét) → hyperlink sang sheet chi tiết HS
- Mỗi học sinh 1 sheet: câu hỏi, đủ đáp án, đáp án đúng, HS chọn
"""
from __future__ import annotations

import os
import re
import sys

from bson import ObjectId
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pymongo import MongoClient

load_dotenv()

QUIZ_ID = "6ab7e539c4110429822a0cb0"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KETQUA_DIR = os.path.join(ROOT, "frontend", "public", "ketqua")

CLASSES = [
    {"file": "Scratch 1.xlsx", "grade_keys": {"Scratch1", "Scratch 1"}, "label": "Scratch1"},
    {"file": "Scratch 2.xlsx", "grade_keys": {"Scratch2", "Scratch 2"}, "label": "Scratch2"},
    {"file": "Scratch 3.xlsx", "grade_keys": {"Scratch3", "Scratch 3"}, "label": "Scratch3"},
    {"file": "Scratch 4.xlsx", "grade_keys": {"Scratch4", "Scratch 4"}, "label": "Scratch4"},
]

HEADER_FILL = PatternFill("solid", fgColor="1E78C8")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
LINK_FONT = Font(color="0563C1", underline="single", bold=True)
THIN = Border(
    left=Side(style="thin", color="B0C4DE"),
    right=Side(style="thin", color="B0C4DE"),
    top=Side(style="thin", color="B0C4DE"),
    bottom=Side(style="thin", color="B0C4DE"),
)


def normalize_name(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def set_internal_sheet_link(cell, sheet_name: str, display: str = "Xem bài làm"):
    """Link nội bộ sheet — tương thích Excel + WPS.

    Dùng công thức HYPERLINK(#...'!A1) thay vì hyperlink file ngoài
    (openpyxl string target khiến WPS báo «Không thể mở tệp»).
    """
    escaped = (sheet_name or "").replace("'", "''")
    # Xóa relationship/target cũ nếu có
    try:
        cell.hyperlink = None
    except Exception:
        pass
    cell.value = f'=HYPERLINK("#\'{escaped}\'!A1","{display}")'
    cell.font = LINK_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")


def sheet_title_for(name: str, used: set) -> str:
    """Tên sheet Excel ≤31 ký tự, không ký tự cấm."""
    base = re.sub(r'[\\/*?:\[\]]', "", (name or "HS").strip()) or "HS"
    base = base[:31]
    title = base
    n = 2
    while title in used:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def mc_score_100(details):
    mc = [d for d in details if d.get("type") == "mcq"]
    earned = sum(int(d.get("points_awarded") or 0) for d in mc)
    max_pts = sum(int(d.get("points") or 0) for d in mc)
    if max_pts <= 0:
        return None, earned, max_pts, mc
    return round(earned / max_pts * 100, 1), earned, max_pts, mc


def format_score(mc100):
    if mc100 is None:
        return "—"
    if float(mc100) == int(mc100):
        return f"{int(mc100)}/100"
    return f"{mc100}/100"


def find_attempt(attempts, name, grade_keys):
    key = normalize_name(name)
    same_grade = []
    any_match = []
    for a in attempts:
        if normalize_name(a.get("student_name") or "") != key:
            continue
        any_match.append(a)
        if (a.get("student_grade") or "").strip() in grade_keys:
            same_grade.append(a)
    pool = same_grade or any_match
    if not pool:
        return None
    # bài nộp mới nhất
    return sorted(
        pool,
        key=lambda x: x.get("submitted_at") or "",
        reverse=True,
    )[0]


def quiz_mc_by_id(quiz):
    out = {}
    for q in quiz.get("questions") or []:
        if q.get("type") != "mcq":
            continue
        out[str(q.get("id"))] = q
    return out


def option_texts(qdoc):
    opts = (qdoc or {}).get("options") or []
    texts = []
    for o in opts:
        if isinstance(o, dict):
            texts.append((o.get("text") or "").strip())
        else:
            texts.append(str(o))
    # pad to 4
    while len(texts) < 4:
        texts.append("")
    return texts[:8]  # allow up to 8


def correct_from_question(qdoc, detail_correct):
    if detail_correct:
        return detail_correct
    for o in (qdoc or {}).get("options") or []:
        if isinstance(o, dict) and o.get("is_correct"):
            return (o.get("text") or "").strip()
    return ""


def read_roster(ws):
    """Đọc danh sách HS từ sheet phiếu (dòng header có 'Họ và tên')."""
    header_row = None
    col_name = col_score = col_comment = None
    for r in range(1, min(ws.max_row or 1, 30) + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        joined = " ".join(str(v) for v in vals if v)
        if "Họ và tên" in joined or "Ho va ten" in joined:
            header_row = r
            for c in range(1, 10):
                v = str(ws.cell(r, c).value or "")
                if "Họ và tên" in v or "tên học sinh" in v.lower():
                    col_name = c
                elif v.strip() in ("Điểm", "Diem"):
                    col_score = c
                elif "Nhận xét" in v or "Nhan xet" in v:
                    col_comment = c
            break
    if not header_row or not col_name:
        raise RuntimeError("Không tìm thấy header danh sách học sinh")

    students = []
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        name = ws.cell(r, col_name).value
        if not name or not str(name).strip():
            # dừng khi gặp ghi chú / footer
            stt = ws.cell(r, 1).value
            if stt is None and not ws.cell(r, col_comment).value:
                continue
            if isinstance(name, str) and name.startswith("Ghi chú"):
                break
            if stt is None:
                break
            continue
        name = str(name).strip()
        if name.startswith("Ghi chú") or name.startswith("©"):
            break
        students.append(
            {
                "row": r,
                "name": name,
                "comment": ws.cell(r, col_comment).value if col_comment else None,
            }
        )
    return {
        "header_row": header_row,
        "col_name": col_name,
        "col_score": col_score or 3,
        "col_comment": col_comment or 4,
        "students": students,
    }


def clear_extra_sheets(wb, keep_title):
    for title in list(wb.sheetnames):
        if title == keep_title:
            continue
        del wb[title]


def write_student_sheet(wb, title, student_name, grade, mc100, earned, max_pts, mc_details, qmap):
    ws = wb.create_sheet(title)
    ws["A1"] = f"CHI TIẾT BÀI LÀM — {student_name}"
    ws["A1"].font = Font(bold=True, size=14, color="1E78C8")
    ws.merge_cells("A1:I1")
    ws["A2"] = f"Lớp: {grade}  |  Điểm trắc nghiệm: {format_score(mc100)}  ({earned}/{max_pts})"
    ws["A2"].font = Font(bold=True, size=11)
    ws.merge_cells("A2:I2")

    headers = [
        "STT",
        "Câu hỏi",
        "Đáp án A",
        "Đáp án B",
        "Đáp án C",
        "Đáp án D",
        "Đáp án đúng",
        "HS chọn",
        "Kết quả",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN

    for i, d in enumerate(mc_details, 1):
        qid = str(d.get("question_id") or "")
        qdoc = qmap.get(qid) or {}
        opts = option_texts(qdoc)
        # nếu thiếu options trong quiz, vẫn ghi từ detail
        correct = correct_from_question(qdoc, d.get("correct_answer"))
        student = (d.get("student_answer") or "").strip() or "(trống)"
        ok = bool(d.get("is_correct"))
        content = (d.get("content") or qdoc.get("content") or "").strip()
        row = 4 + i
        values = [
            i,
            content,
            opts[0] if len(opts) > 0 else "",
            opts[1] if len(opts) > 1 else "",
            opts[2] if len(opts) > 2 else "",
            opts[3] if len(opts) > 3 else "",
            correct,
            student,
            "Đúng" if ok else "Sai",
        ]
        for c, val in enumerate(values, 1):
            cell = ws.cell(row, c, val)
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        fill = OK_FILL if ok else BAD_FILL
        ws.cell(row, 9).fill = fill
        ws.cell(row, 9).font = Font(bold=True, color="006100" if ok else "9C0006")

    widths = {
        "A": 6,
        "B": 48,
        "C": 28,
        "D": 28,
        "E": 28,
        "F": 28,
        "G": 28,
        "H": 28,
        "I": 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 22
    for r in range(5, 5 + len(mc_details)):
        ws.row_dimensions[r].height = 60
    ws.freeze_panes = "A5"
    return ws


def process_class(wb_path, grade_keys, label, attempts, qmap):
    wb = load_workbook(wb_path)
    main = wb.active
    main_title = main.title
    roster = read_roster(main)
    header_row = roster["header_row"]
    col_score = roster["col_score"]
    col_comment = roster["col_comment"]
    col_link = col_comment + 1

    # Xóa sheet phụ cũ (Chi_tiet_TN, Anh_TN, sheet HS cũ…)
    clear_extra_sheets(wb, main_title)
    main = wb[main_title]

    # Header cột Xem bài làm
    main.cell(header_row, col_link).value = "Xem bài làm"
    main.cell(header_row, col_link).font = Font(bold=True)
    main.cell(header_row, col_link).alignment = Alignment(horizontal="center")

    # Dọn cột thừa bên phải link (từ lần xuất trước)
    for c in range(col_link + 1, col_link + 6):
        if main.cell(header_row, c).value in (
            "Điểm TN (/100)",
            "Link xem dự án (.sb3)",
            "Link kết quả / ảnh TN",
            "Link xem dự án",
            "Link kết quả",
        ):
            for r in range(header_row, header_row + 20):
                main.cell(r, c).value = None
                main.cell(r, c).hyperlink = None

    used_titles = {main_title}
    ok_count = 0
    miss = []

    for st in roster["students"]:
        name = st["name"]
        row = st["row"]
        attempt = find_attempt(attempts, name, grade_keys)
        if not attempt:
            main.cell(row, col_score).value = "—"
            main.cell(row, col_link).value = "(chưa có bài)"
            main.cell(row, col_link).hyperlink = None
            miss.append(name)
            continue

        details = attempt.get("details") or []
        mc100, earned, max_pts, mc = mc_score_100(details)
        grade = (attempt.get("student_grade") or label).strip()

        # Điểm trên phiếu = chỉ TN
        main.cell(row, col_score).value = format_score(mc100)

        title = sheet_title_for(name, used_titles)
        write_student_sheet(wb, title, name, grade, mc100, earned, max_pts, mc, qmap)

        set_internal_sheet_link(main.cell(row, col_link), title, "Xem bài làm")
        ok_count += 1

    main.column_dimensions[get_column_letter(col_link)].width = 16

    # Cập nhật ghi chú: điểm trên phiếu = trắc nghiệm
    for r in range(1, (main.max_row or 1) + 1):
        v = main.cell(r, 1).value
        if isinstance(v, str) and v.startswith("Ghi chú:"):
            main.cell(r, 1).value = (
                "Ghi chú: Điểm trên phiếu là điểm trắc nghiệm (thang 100). "
                "Bấm «Xem bài làm» để mở sheet chi tiết từng câu (đáp án đúng / đáp án học sinh chọn). "
                "Phụ huynh vui lòng liên hệ giáo viên nếu cần trao đổi thêm về lộ trình học tập của em."
            )
            break

    wb.save(wb_path)
    return ok_count, miss, len(roster["students"])


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    client = MongoClient(os.getenv("MONGODB_URI"))
    db = client[os.getenv("MONGODB_DB_NAME")]
    quiz = db.quizzes.find_one({"_id": ObjectId(QUIZ_ID)}) or db.quizzes.find_one({"_id": QUIZ_ID})
    if not quiz:
        print("QUIZ_NOT_FOUND")
        sys.exit(1)

    attempts = list(db.quiz_attempts.find({"quiz_id": str(quiz["_id"])}))
    qmap = quiz_mc_by_id(quiz)
    print(f"quiz={quiz.get('title')} attempts={len(attempts)} mc_questions={len(qmap)}")

    for cfg in CLASSES:
        path = os.path.join(KETQUA_DIR, cfg["file"])
        if not os.path.exists(path):
            print("MISSING_FILE", path)
            continue
        ok, miss, total = process_class(path, cfg["grade_keys"], cfg["label"], attempts, qmap)
        print(f"OK {cfg['file']}: matched={ok}/{total} miss={miss}")


if __name__ == "__main__":
    main()
