"""Quản lý báo cáo dạy học dạng Excel trong trang admin."""

from __future__ import annotations

import os
from pathlib import Path

from flask import Blueprint, current_app, jsonify, request, send_file
from flask_jwt_extended import get_jwt, jwt_required
from openpyxl import load_workbook


reports_bp = Blueprint("reports", __name__, url_prefix="/api/reports")

ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}
MAX_ROWS = 2000
MAX_COLUMNS = 200


def _admin_required():
    return get_jwt().get("role") == "admin"


def _reports_root() -> Path:
    configured = os.getenv("TEACHING_REPORTS_FOLDER")
    if configured:
        root = Path(configured)
    else:
        root = Path(current_app.root_path).parent / "frontend" / "public" / "Baocaodayhoc"
    root.mkdir(parents=True, exist_ok=True)
    return root.resolve()


def _safe_file(relative_path: str) -> Path:
    value = (relative_path or "").replace("\\", "/").strip("/")
    if not value:
        raise ValueError("Thiếu đường dẫn file")

    root = _reports_root()
    candidate = (root / value).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("Đường dẫn file không hợp lệ") from exc

    if candidate.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ValueError("Chỉ hỗ trợ file .xlsx và .xlsm")
    return candidate


def _cell_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


@reports_bp.route("", methods=["GET"])
@jwt_required()
def list_reports():
    if not _admin_required():
        return jsonify({"message": "Không có quyền truy cập"}), 403

    root = _reports_root()
    grouped = {}
    for path in sorted(root.rglob("*"), key=lambda p: str(p).lower()):
        if not path.is_file() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        relative = path.relative_to(root)
        class_name = relative.parts[0] if len(relative.parts) > 1 else "Chưa phân lớp"
        stat = path.stat()
        grouped.setdefault(class_name, []).append({
            "name": path.name,
            "path": relative.as_posix(),
            "size": stat.st_size,
            "updated_at": stat.st_mtime,
        })

    classes = [
        {"name": name, "files": files}
        for name, files in sorted(grouped.items(), key=lambda item: item[0].lower())
    ]
    return jsonify({"classes": classes, "file_count": sum(len(c["files"]) for c in classes)})


@reports_bp.route("/file", methods=["GET"])
@jwt_required()
def read_report():
    if not _admin_required():
        return jsonify({"message": "Không có quyền truy cập"}), 403

    try:
        path = _safe_file(request.args.get("path"))
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400
    if not path.is_file():
        return jsonify({"message": "Không tìm thấy file báo cáo"}), 404

    try:
        workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
        sheets = []
        for sheet in workbook.worksheets:
            row_count = min(sheet.max_row or 1, MAX_ROWS)
            column_count = min(sheet.max_column or 1, MAX_COLUMNS)
            rows = [
                [_cell_value(sheet.cell(row=row, column=column).value)
                 for column in range(1, column_count + 1)]
                for row in range(1, row_count + 1)
            ]
            sheets.append({
                "name": sheet.title,
                "rows": rows,
                "row_count": row_count,
                "column_count": column_count,
                "truncated": (sheet.max_row or 1) > MAX_ROWS or (sheet.max_column or 1) > MAX_COLUMNS,
            })
        workbook.close()
        return jsonify({"name": path.name, "path": request.args.get("path"), "sheets": sheets})
    except Exception as exc:
        current_app.logger.exception("Không đọc được Excel %s", path)
        return jsonify({"message": f"Không đọc được file Excel: {exc}"}), 400


@reports_bp.route("/file", methods=["PUT"])
@jwt_required()
def update_report():
    if not _admin_required():
        return jsonify({"message": "Không có quyền truy cập"}), 403

    data = request.get_json(silent=True) or {}
    try:
        path = _safe_file(data.get("path"))
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400
    if not path.is_file():
        return jsonify({"message": "Không tìm thấy file báo cáo"}), 404

    changes = data.get("changes") or []
    if not isinstance(changes, list) or len(changes) > 10000:
        return jsonify({"message": "Danh sách thay đổi không hợp lệ hoặc quá lớn"}), 400

    try:
        workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
        changed = 0
        for item in changes:
            if not isinstance(item, dict):
                continue
            sheet_name = str(item.get("sheet") or "")
            row = int(item.get("row") or 0)
            column = int(item.get("column") or 0)
            if sheet_name not in workbook.sheetnames:
                continue
            if not (1 <= row <= MAX_ROWS and 1 <= column <= MAX_COLUMNS):
                continue
            workbook[sheet_name].cell(row=row, column=column).value = item.get("value")
            changed += 1
        workbook.save(path)
        workbook.close()
        return jsonify({"message": "Đã lưu báo cáo", "changed": changed})
    except Exception as exc:
        current_app.logger.exception("Không lưu được Excel %s", path)
        return jsonify({"message": f"Không lưu được file Excel: {exc}"}), 400


@reports_bp.route("/download", methods=["GET"])
@jwt_required()
def download_report():
    if not _admin_required():
        return jsonify({"message": "Không có quyền truy cập"}), 403

    try:
        path = _safe_file(request.args.get("path"))
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400
    if not path.is_file():
        return jsonify({"message": "Không tìm thấy file báo cáo"}), 404
    return send_file(path, as_attachment=True, download_name=path.name)
